
from pathlib import Path
from datetime import datetime

from django.test.runner import DiscoverRunner
from docx import Document


class WordReportRunner(DiscoverRunner):
    def suite_result(self, suite, result, **kwargs):
        failures_count = super().suite_result(
            suite, result, **kwargs
        )

        project_dir = Path(__file__).resolve().parent.parent
        report_path = project_dir / "Django_blog_lab_tests.docx"

        if report_path.exists():
            doc = Document(report_path)
        else:
            doc = Document()
            doc.add_heading(
                "Отчёт о тестировании Django-проекта",
                level=1
            )

        doc.add_heading(
            f"Результат запуска — {datetime.now():%d.%m.%Y %H:%M:%S}",
            level=2
        )

        doc.add_paragraph(
            f"Всего выполнено тестов: {result.testsRun}"
        )
        doc.add_paragraph(
            f"Ошибок: {len(result.errors)}"
        )
        doc.add_paragraph(
            f"Проваленных тестов: {len(result.failures)}"
        )
        doc.add_paragraph(
            f"Пропущено тестов: {len(result.skipped)}"
        )

        if failures_count == 0:
            doc.add_paragraph(
                "Итог: OK — все запущенные тесты прошли успешно."
            )
        else:
            doc.add_paragraph(
                "Итог: обнаружены ошибки. Нужно исправить тесты."
            )

        for test, details in result.failures:
            doc.add_heading(f"FAIL: {test}", level=3)
            doc.add_paragraph(details)

        for test, details in result.errors:
            doc.add_heading(f"ERROR: {test}", level=3)
            doc.add_paragraph(details)

        doc.save(report_path)
        return failures_count
